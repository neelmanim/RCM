"""v13.11.0 — meeting record, status-log user id, lifecycle labels, Excel export, per-SDR meetings fallback."""
import io
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest

import meeting_log
import migrations
import models
from conftest import (create_test_call, create_test_lead, create_test_pod, create_test_user, grant_export)
from routes import analytics_routes
from services import lead_lifecycle as lc


def _meetings(db, lead_id):
    db.expire_all()
    return db.query(models.LeadMeeting).filter(models.LeadMeeting.lead_id == lead_id).order_by(models.LeadMeeting.booked_at).all()


# ── meeting record, written by the one shared status-log helper ─────────────────────────────────────────────────
def test_meeting_stage_list_matches_analytics():
    assert list(meeting_log.MEETING_REACHED) == analytics_routes.MEETING_REACHED_STATUSES


def test_booking_is_recorded_with_who_and_stamps_lead_closed_at(db):
    sdr = create_test_user(db, email="s@t.com", id="sdr-1", name="Sam SDR")
    lead = create_test_lead(db, email="m1@t.com", status="Calling")
    lead.status = "Meeting Scheduled"
    models.log_status_change(db, lead.id, "Calling", "Meeting Scheduled", "Sam SDR", changed_by_user_id=sdr.id)
    db.commit()
    [m] = _meetings(db, lead.id)
    assert (m.status, m.booked_by_user_id, m.source) == ("booked", "sdr-1", "app")
    db.refresh(lead)
    assert lead.lead_closed_at is not None                       # the call-outcome / Aircall paths never stamped this
    log = db.query(models.LeadStatusLog).filter_by(lead_id=lead.id).one()
    assert log.changed_by_user_id == "sdr-1"


def test_same_status_resave_and_second_hook_do_not_double_book(db):
    lead = create_test_lead(db, email="m2@t.com", status="Meeting Scheduled")
    models.log_status_change(db, lead.id, "Meeting Scheduled", "Meeting Scheduled", "x")
    models.log_status_change(db, lead.id, "Calling", "Meeting Scheduled", "x")
    models.log_status_change(db, lead.id, "Meeting Scheduled", "1st Discovery Meeting", "x")
    db.commit()
    assert [m.status for m in _meetings(db, lead.id)] == ["held"]       # one booking, later marked held — never two


def test_discovery_after_booking_marks_it_held_and_skipping_the_booking_creates_a_held_row(db):
    a = create_test_lead(db, email="m3@t.com", status="Meeting Scheduled")
    models.log_status_change(db, a.id, "Calling", "Meeting Scheduled", "x")
    models.log_status_change(db, a.id, "Meeting Scheduled", "1st Discovery Meeting", "x")
    b = create_test_lead(db, email="m4@t.com", status="1st Discovery Meeting")
    models.log_status_change(db, b.id, "Calling", "1st Discovery Meeting", "x")          # never booked: the "skip" path
    models.log_status_change(db, b.id, "1st Discovery Meeting", "Discovery Complete", "x")  # later stage: no second row
    db.commit()
    assert [m.status for m in _meetings(db, a.id)] == ["held"]
    assert [m.status for m in _meetings(db, b.id)] == ["held"]


def test_no_show_then_rebook_is_two_bookings_and_disqualify_cancels(db):
    lead = create_test_lead(db, email="m5@t.com", status="Calling")
    models.log_status_change(db, lead.id, "Calling", "Meeting Scheduled", "x")
    meeting_log.mark_no_show(db, lead)
    models.log_status_change(db, lead.id, "Meeting Scheduled", "Calling", "x")
    models.log_status_change(db, lead.id, "Calling", "Meeting Scheduled", "x")
    models.log_status_change(db, lead.id, "Meeting Scheduled", "Disqualified", "x")
    db.commit()
    assert [m.status for m in _meetings(db, lead.id)] == ["no_show", "cancelled"]


def test_source_is_read_from_the_actor_string(db):
    a = create_test_lead(db, email="m6@t.com")
    b = create_test_lead(db, email="m7@t.com")
    models.log_status_change(db, a.id, "Calling", "Meeting Scheduled", "Aircall tag: Meeting")
    models.log_status_change(db, b.id, None, "Meeting Scheduled", "Salesforce sync")
    db.commit()
    assert _meetings(db, a.id)[0].source == "aircall-tag" and _meetings(db, b.id)[0].source == "salesforce-sync"


def test_status_update_endpoint_books_a_meeting_and_records_the_user(client, db):
    lead = create_test_lead(db, email="m8@t.com", status="Calling")
    create_test_call(db, lead.id, None)
    r = client.patch(f"/api/leads/{lead.id}", json={"status": "Meeting Scheduled"})
    assert r.status_code == 200, r.text
    [m] = _meetings(db, lead.id)
    assert m.status == "booked" and m.booked_by_user_id == "test-user-id"


def test_no_show_endpoint_marks_the_open_booking(client, db):
    lead = create_test_lead(db, email="m9@t.com", status="Calling")
    models.log_status_change(db, lead.id, "Calling", "Meeting Scheduled", "x")
    lead.status = "Meeting Scheduled"
    db.commit()
    r = client.post(f"/api/leads/{lead.id}/no-show", json={"reason": "did not join the call at all"})
    assert r.status_code == 200, r.text
    assert [m.status for m in _meetings(db, lead.id)] == ["no_show"]


# ── backfill ────────────────────────────────────────────────────────────────────────────────────────────────────
def test_backfill_creates_missing_records_stamps_closed_at_and_is_idempotent(db, engine):
    sam = create_test_user(db, email="sam@t.com", id="sam", name="Sam Seller")
    booked = create_test_lead(db, email="b1@t.com", status="Meeting Scheduled")
    held = create_test_lead(db, email="b2@t.com", status="Demo Done")
    for l in (booked, held):
        l.status_changed_at = datetime(2026, 9, 1, tzinfo=timezone.utc)
        l.lead_closed_at = None
    db.add(models.LeadStatusLog(lead_id=booked.id, from_status="Calling", to_status="Meeting Scheduled", changed_by="Sam Seller",
                                changed_at=datetime(2026, 8, 30, tzinfo=timezone.utc)))
    db.commit()
    migrations._backfill_lead_meetings(engine)
    [mb] = _meetings(db, booked.id)
    [mh] = _meetings(db, held.id)
    assert (mb.status, mb.booked_by_user_id, mb.source) == ("booked", "sam", "backfill")
    assert (mh.status, mh.booked_by_user_id) == ("held", None)
    db.refresh(booked)
    assert booked.lead_closed_at is not None
    migrations._backfill_lead_meetings(engine)
    assert len(_meetings(db, booked.id)) == 1


# ── lifecycle labels ────────────────────────────────────────────────────────────────────────────────────────────
def _own(db, user, lead, at=None):
    db.execute(models.lead_assignments.insert().values(user_id=user.id, lead_id=lead.id,
                                                       assigned_at=at or datetime.now(timezone.utc) - timedelta(days=2)))
    db.commit()


def test_lifecycle_states(db):
    old = create_test_user(db, email="old@t.com", id="old", name="Old Owner")
    me = create_test_user(db, email="me@t.com", id="me", name="Me Rep")
    unassigned = create_test_lead(db, email="u0@t.com", status="Calling")
    fresh = create_test_lead(db, email="u1@t.com", status="Lead Assigned")
    inherited = create_test_lead(db, email="u2@t.com", status="Calling")
    working = create_test_lead(db, email="u3@t.com", status="Calling")
    booked = create_test_lead(db, email="u4@t.com", status="Meeting Scheduled")
    closed = create_test_lead(db, email="u5@t.com", status="Disqualified")
    parked = create_test_lead(db, email="u6@t.com", status="No Phone - Parked")
    other = create_test_lead(db, email="u7@t.com", status="Meeting Booked")
    for l in (fresh, inherited, working, booked):
        _own(db, me, l)
    create_test_call(db, inherited.id, old.id, called_at=datetime.now(timezone.utc) - timedelta(days=10))   # worked by someone else
    create_test_call(db, working.id, me.id, called_at=datetime.now(timezone.utc) - timedelta(hours=3))
    got = lc.states_for(db, [unassigned, fresh, inherited, working, booked, closed, parked, other])
    assert got == {unassigned.id: "Unassigned", fresh.id: "New to me", inherited.id: "Inherited",
                   working.id: "In progress", booked.id: "Meeting booked", closed.id: "Closed",
                   parked.id: "On hold", other.id: "Other"}
    assert lc.summary_for_user(db, "me") == {"New to me": 1, "Inherited": 1, "In progress": 1, "Meeting booked": 1,
                                             "Closed": 0, "On hold": 0, "Other": 0}


def test_reset_with_the_same_owner_reads_new_but_a_new_owner_reads_inherited(db):
    me = create_test_user(db, email="me@t.com", id="me", name="Me Rep")
    lead = create_test_lead(db, email="r1@t.com", status="Lead Assigned")
    _own(db, me, lead, datetime.now(timezone.utc) - timedelta(days=9))
    create_test_call(db, lead.id, me.id, called_at=datetime.now(timezone.utc) - timedelta(days=5))
    db.add(models.LeadStatusLog(lead_id=lead.id, from_status="Calling", to_status="Lead Assigned", changed_by="Me Rep",
                                changed_at=datetime.now(timezone.utc) - timedelta(days=1)))
    db.commit()
    assert lc.states_for(db, [lead])[lead.id] == "New to me"
    other = create_test_lead(db, email="r2@t.com", status="Lead Assigned")
    prev = create_test_user(db, email="prev@t.com", id="prev", name="Prev Owner")
    _own(db, me, other)
    create_test_call(db, other.id, prev.id, called_at=datetime.now(timezone.utc) - timedelta(days=5))
    assert lc.states_for(db, [other])[other.id] == "Inherited"


def test_lifecycle_endpoint_for_the_caller(client, db):
    create_test_user(db, email="admin@test.com", id="test-user-id", name="Test Admin")
    lead = create_test_lead(db, email="e1@t.com", status="Calling")
    _own(db, db.get(models.User, "test-user-id"), lead)
    r = client.get("/api/my-day/lifecycle")
    assert r.status_code == 200 and r.json()["counts"]["New to me"] == 1


# ── Excel export ────────────────────────────────────────────────────────────────────────────────────────────────
def _export(client, **params):
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    r = client.get("/api/exports/activity.xlsx", params={"date_from": today, "date_to": today, **params})
    return r


def test_export_has_calls_and_meetings_sheets_and_spreadsheet_safe_cells(client, db):
    sam = create_test_user(db, email="sam@t.com", id="sam", name="=Sam")
    lead = create_test_lead(db, email="x1@t.com", first_name="Ann", last_name="Lee", status="Meeting Scheduled")
    lead.company = "Acme"
    create_test_call(db, lead.id, "sam", outcome="Meeting Scheduled")
    models.log_status_change(db, lead.id, "Calling", "Meeting Scheduled", "x", changed_by_user_id="sam")
    db.commit()
    r = _export(client)
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Calls", "Meetings", "Read me"]
    calls = list(wb["Calls"].iter_rows(values_only=True))
    assert calls[1][1] == "'=Sam" and calls[1][2] == "Ann Lee" and calls[1][3] == "Acme" and calls[1][12] == "Yes"
    meets = list(wb["Meetings"].iter_rows(values_only=True))
    assert meets[1][1] == "'=Sam" and meets[1][7] == "booked"


def test_export_is_admin_only_and_limited_in_range(client, client_as_sdr, db):
    assert _export(client_as_sdr).status_code == 403
    r = client.get("/api/exports/activity.xlsx", params={"date_from": "2026-01-01", "date_to": "2026-12-31"})
    assert r.status_code == 400
    r = client.get("/api/exports/activity.xlsx", params={"date_from": "2026-02-01", "date_to": "2026-01-01"})
    assert r.status_code == 400


def test_pod_admin_export_only_includes_their_pod(client_as_pod_admin, db):
    pod = create_test_pod(db, name="Mine")
    other_pod = create_test_pod(db, name="Theirs")
    pod.id = "test-pod-id"
    db.commit()
    mine = create_test_user(db, email="mine@t.com", id="mine", name="Mine Rep", pod_id="test-pod-id")
    theirs = create_test_user(db, email="theirs@t.com", id="theirs", name="Their Rep", pod_id=other_pod.id)
    lead = create_test_lead(db, email="x2@t.com")
    create_test_call(db, lead.id, "mine")
    create_test_call(db, lead.id, "theirs")
    grant_export(db, "pod-admin-id", role="Pod Admin")
    wb = openpyxl.load_workbook(io.BytesIO(_export(client_as_pod_admin).content))
    assert [r[1] for r in list(wb["Calls"].iter_rows(values_only=True))[1:]] == ["Mine Rep"]


# ── per-SDR Analytics meetings falls back like the trend does ───────────────────────────────────────────────────
def test_per_sdr_meetings_credit_the_booker_not_the_current_owner(client, db):
    booker = create_test_user(db, email="p@t.com", id="p-sdr", name="Pat SDR")
    holder = create_test_user(db, email="h2@t.com", id="h-sdr", name="Hal SDR")
    lead = create_test_lead(db, email="a1@t.com", status="Meeting Scheduled")
    models.log_status_change(db, lead.id, "Calling", "Meeting Scheduled", "Pat", changed_by_user_id="p-sdr")
    _own(db, holder, lead)                                   # handed on afterwards
    db.commit()
    rows = {x["sdr_id"]: x for x in client.get("/api/admin/analytics/sdr-table", params={"preset": "7d"}).json()["sdrs"]}
    assert rows["p-sdr"]["meetings"] == 1 and rows["h-sdr"]["meetings"] == 0


# ── cross-screen reconciliation: two screens must agree for the same person and day ─────────────────────────────
def test_export_call_rows_equal_my_day_calls(client, db):
    from routes import my_day_routes
    sam = create_test_user(db, email="rc@t.com", id="rc", name="Rec Rep")
    lead = create_test_lead(db, email="rc1@t.com")
    create_test_call(db, lead.id, "rc", outcome="No Answer")
    create_test_call(db, lead.id, "rc", outcome="Interested")
    from conftest import create_test_dialer_call
    create_test_dialer_call(db, lead.id, "rc", direction="outbound")
    mine = my_day_routes._build(db, {"sub": "rc"}, None, 0)["counts"]["calls"]
    wb = openpyxl.load_workbook(io.BytesIO(_export(client, sdr_id="rc").content))
    assert len(list(wb["Calls"].iter_rows(values_only=True))) - 1 == mine == 3


# ── offboarding: one removal that really removes ────────────────────────────────────────────────────────────────
def _leaver(db, **kw):
    import access_db
    u = create_test_user(db, email="gone@t.com", id="gone", name="Gone Rep", **kw)
    u.dialer_enabled = u.sms_enabled = u.whatsapp_enabled = u.email_sync_enabled = True
    db.add(models.UserMailbox(user_id="gone", email_address="gone@t.com", nylas_grant_id="g", status="connected"))
    access_db.add_allowed_user(db, "gone@t.com", "Gone Rep", "SDR", "admin")
    lead = create_test_lead(db, email="gl@t.com", status="Calling")
    _own(db, u, lead)
    db.commit()
    return u, lead


def test_delete_offboards_instead_of_deleting(client, db):
    u, lead = _leaver(db)
    r = client.delete("/api/admin/users/gone")
    assert r.status_code == 200, r.text
    db.expire_all()
    u = db.get(models.User, "gone")
    assert u is not None and u.left_at is not None                                   # row kept: history needs a name
    assert not (u.dialer_enabled or u.sms_enabled or u.whatsapp_enabled or u.email_sync_enabled)
    assert db.query(models.UserMailbox).filter_by(user_id="gone").one().status == "disconnected"
    assert db.query(models.lead_assignments).filter_by(user_id="gone").count() == 0  # portfolio released
    assert db.query(models.AllowedUser).filter_by(email="gone@t.com").count() == 0   # cannot sign in


def test_a_leaver_is_gone_from_settings_and_marked_left_elsewhere(client, db):
    _leaver(db)
    client.delete("/api/admin/users/gone")
    assert "gone" not in [x["id"] for x in client.get("/api/admin/dialer/sdr-configs").json()]
    assert "gone" not in [x["id"] for x in client.get("/api/admin/users").json()]
    row = [x for x in client.get("/api/admin/users", params={"include_left": True}).json() if x["id"] == "gone"][0]
    assert row["left"] is True and row["left_at"]
    sdr_rows = client.get("/api/admin/analytics/sdr-table").json()["sdrs"]
    assert [x["left"] for x in sdr_rows if x["sdr_id"] == "gone"] == [True]


def test_someone_removed_by_the_old_csv_path_still_reads_as_left(client, db):
    import access_db
    create_test_user(db, email="old@t.com", id="oldcsv", name="Old CSV")     # no allowed_users entry = cannot sign in
    from services import offboarding
    assert "oldcsv" in offboarding.left_ids(db)


def test_adding_a_leaver_back_reactivates_the_same_row(client, db):
    _leaver(db)
    client.delete("/api/admin/users/gone")
    r = client.post("/api/admin/users", json={"email": "gone@t.com", "name": "Gone Rep", "role": "SDR"})
    assert r.status_code == 200 and "re-activated" in r.json()["message"]
    db.expire_all()
    assert db.get(models.User, "gone").left_at is None
    assert db.query(models.AllowedUser).filter_by(email="gone@t.com").count() == 1


@pytest.fixture(autouse=True)
def _fresh_response_cache():
    """Postgres runs have the response cache on; a leaderboard cached by one test must not leak into the next."""
    import cache
    cache.clear_all()


# ── leaderboard = meetings BOOKED, not leads currently held ─────────────────────────────────────────────────────
def test_leaderboard_counts_bookings_not_inherited_stock(client, db):
    booker = create_test_user(db, email="b@t.com", id="booker", name="Booker")
    holder = create_test_user(db, email="h@t.com", id="holder", name="Holder")
    import access_db
    for u in (booker, holder):
        access_db.add_allowed_user(db, u.email, u.name, "SDR", "admin")
    lead = create_test_lead(db, email="lb@t.com", status="Calling")
    models.log_status_change(db, lead.id, "Calling", "Meeting Scheduled", "Booker", changed_by_user_id="booker")
    lead.status = "Meeting Scheduled"
    db.commit()
    _own(db, holder, lead)                                  # the lead is later handed to someone else
    rows = {r["id"]: r for r in client.get("/api/leaderboard").json()}
    assert rows["booker"]["meetings_scheduled"] == 1
    assert rows["holder"]["meetings_scheduled"] == 0 and rows["holder"]["meetings_in_stage_now"] == 1


def test_leaderboard_marks_a_leaver_who_still_has_activity_and_hides_one_who_has_none(client, db):
    u, _ = _leaver(db)
    other = create_test_lead(db, email="lb2@t.com", status="Calling")
    create_test_call(db, other.id, "gone")
    client.delete("/api/admin/users/gone")
    assert [r["left"] for r in client.get("/api/leaderboard").json() if r["id"] == "gone"] == [True]
    db.query(models.CallLog).delete()
    db.commit()
    import cache
    cache.clear_all()
    assert "gone" not in [r["id"] for r in client.get("/api/leaderboard").json()]


# ── data quality: the backfill is readable without database access ──────────────────────────────────────────────
def test_meetings_data_quality_endpoint(client, db):
    lead = create_test_lead(db, email="dq@t.com", status="Meeting Scheduled")
    lead.lead_closed_at = None
    db.commit()
    d = client.get("/api/admin/data-quality/meetings").json()
    assert d["meeting_stage_leads"] == 1 and d["meeting_stage_leads_without_a_booking_record"] == 1
    assert d["meeting_stage_leads_without_lead_closed_at"] == 1
    migrations._backfill_lead_meetings(db.get_bind())
    d = client.get("/api/admin/data-quality/meetings").json()
    assert d["meeting_stage_leads_without_a_booking_record"] == 0 and d["by_source"] == {"backfill": 1}
