"""Funnel report (services/funnel.py + /api/admin/analytics/sales-funnel): ownership-start accounts, conversations, bookings by
booker, SAL, held / no-show, campaign grouping, account key, and who may see which pod."""
from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient

import access_db
import models
from services import funnel
from tests.conftest import (_build_test_app, _make_user_payload, create_test_booking, create_test_call, create_test_lead,
                            create_test_user, create_test_pod, ensure_pods)

D = datetime(2026, 2, 10, 10, 0, tzinfo=timezone.utc)
Q = "?from=2026-02-09&to=2026-02-11"


def _as(db, payload):
    from auth import get_current_user, require_admin
    from database import get_db, get_ai_db
    app = _build_test_app()
    app.dependency_overrides[get_db] = app.dependency_overrides[get_ai_db] = lambda: db
    app.dependency_overrides[get_current_user] = app.dependency_overrides[require_admin] = lambda: payload
    return TestClient(app)


def _super(db):
    return _as(db, _make_user_payload("Super Admin", "sa", "sa@t.com", "SA"))


def _own(db, lead, user_id, at=D):
    db.add(models.LeadOwnershipPeriod(lead_id=lead.id, user_id=user_id, valid_from=at, reason="test"))
    db.commit()


def _tag(db, lead, name):
    tag = db.query(models.Tag).filter_by(name=name).first() or models.Tag(name=name)
    db.add(tag)
    db.flush()
    db.execute(models.lead_tags.insert().values(lead_id=lead.id, tag_id=tag.id))
    db.commit()


# ── account key ─────────────────────────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("company,website,email,expected", [
    ("Acme Inc", "https://www.acme.com/x", None, "acme.com"),       # website domain wins
    ("Acme Inc", None, "a@acme.io", "acme.io"),                     # then a company e-mail domain
    ("Acme, Inc.", None, "a@gmail.com", "acme"),                    # public mail is not an account
    ("ACME  Ltd", None, None, "acme"),                              # then the company name, case/suffix ignored
    (None, None, None, None),
])
def test_account_key(company, website, email, expected):
    assert funnel.account_key(company, website, email) == expected


# ── the funnel numbers ──────────────────────────────────────────────────────────────────────────────────────────
@pytest.fixture
def team(db):
    ensure_pods(db, "p1")
    create_test_user(db, email="a@t.com", id="a", name="Asha", role="SDR", pod_id="p1")
    create_test_user(db, email="b@t.com", id="b", name="Bala", role="SDR", pod_id="p1")
    return db


def test_counts_for_the_rep_in_the_window_only(team):
    db = team
    l1 = create_test_lead(db, email="x@acme.com", company="Acme", sf_lead_id="00Q" + "A" * 15, pod_id="p1")
    l2 = create_test_lead(db, email="y@acme.com", company="Acme Inc", pod_id="p1")
    l3 = create_test_lead(db, email="z@other.com", company="Other", pod_id="p1")
    for lead in (l1, l2, l3):
        _own(db, lead, "a")
    _own(db, l3, "b", datetime(2026, 1, 1, tzinfo=timezone.utc))     # outside the window
    create_test_call(db, l1.id, "a", outcome="Interested", called_at=D)
    create_test_call(db, l1.id, "a", outcome="No Answer", called_at=D)          # not a conversation
    m = create_test_booking(db, l1, "a", booked_at=D)
    create_test_booking(db, l2, "a", booked_at=D)
    m.status, m.held_at = "held", D
    db.commit()
    r = _super(db).get("/api/admin/analytics/sales-funnel" + Q + "&group_by=sdr").json()
    rows = {x["key"]: x for x in r["rows"]}
    assert rows["a"]["contacts"] == 3 and rows["a"]["accounts"] == 2      # x@ and y@acme.com share a domain = one account; other.com is the second
    assert rows["a"]["conversations"] == 1
    assert rows["a"]["booked"] == 2 and rows["a"]["sal"] == 1 and rows["a"]["held"] == 1
    assert rows["a"]["held_rate"] == 50.0
    assert "b" not in rows or rows["b"]["contacts"] == 0               # the old ownership period is outside the window
    assert r["total"]["booked"] == 2


def test_cancelled_booking_and_no_show(team):
    db = team
    lead = create_test_lead(db, email="n@t.com", company="N", pod_id="p1")
    create_test_booking(db, lead, "a", status="cancelled", booked_at=D)
    ns = create_test_booking(db, lead, "a", booked_at=D)
    ns.status, ns.no_show_at = "no_show", D
    db.commit()
    t = _super(db).get("/api/admin/analytics/sales-funnel" + Q).json()["total"]
    assert t["booked"] == 1 and t["no_show"] == 1 and t["no_show_rate"] == 100.0


def test_campaign_grouping_and_filter_never_hide_untagged(team):
    db = team
    tagged = create_test_lead(db, email="t@t.com", company="T", pod_id="p1")
    bare = create_test_lead(db, email="u@u.com", company="U", pod_id="p1")
    _tag(db, tagged, "US-Partner-FY27Q3")
    create_test_booking(db, tagged, "a", booked_at=D)
    create_test_booking(db, bare, "a", booked_at=D)
    c = _super(db)
    rows = {x["label"]: x for x in c.get("/api/admin/analytics/sales-funnel" + Q + "&group_by=campaign").json()["rows"]}
    assert rows["US-Partner-FY27Q3"]["booked"] == 1 and rows["No campaign"]["booked"] == 1
    only = c.get("/api/admin/analytics/sales-funnel" + Q + "&campaign=No campaign").json()["total"]
    assert only["booked"] == 1


def test_deleted_and_test_leads_are_not_counted(team):
    db = team
    gone = create_test_lead(db, email="g@t.com", company="G", pod_id="p1")
    gone.deleted_at = D
    fake = create_test_lead(db, email="f@t.com", company="F", pod_id="p1")
    fake.is_test = True
    db.commit()
    for lead in (gone, fake):
        create_test_booking(db, lead, "a", booked_at=D)
    assert _super(db).get("/api/admin/analytics/sales-funnel" + Q).json()["total"]["booked"] == 0


# ── who may see what, validation, export ────────────────────────────────────────────────────────────────────────
def test_pod_admin_sees_only_their_own_pod_and_podless_sees_nothing(db):
    ensure_pods(db, "p1", "p2")
    create_test_user(db, email="a@t.com", id="a", role="SDR", pod_id="p1")
    create_test_user(db, email="c@t.com", id="c", role="SDR", pod_id="p2")
    for uid, pod in (("a", "p1"), ("c", "p2")):
        create_test_booking(db, create_test_lead(db, email=f"{uid}@l.com", company=uid, pod_id=pod), uid, booked_at=D)
    mine = _as(db, _make_user_payload("Pod Admin", "pa", "pa@t.com", "PA", pod_id="p1"))
    assert mine.get("/api/admin/analytics/sales-funnel" + Q + "&pod_id=p2").json()["total"]["booked"] == 1    # pod_id param ignored
    assert mine.get("/api/admin/analytics/sales-funnel" + Q + "&group_by=sdr").json()["rows"][0]["key"] == "a"
    none = _as(db, _make_user_payload("Pod Admin", "pb", "pb@t.com", "PB", pod_id=access_db.NO_POD))
    assert none.get("/api/admin/analytics/sales-funnel" + Q).json()["total"]["booked"] == 0


def test_sdr_is_refused_and_bad_input_is_a_400(db, client_as_sdr):
    assert client_as_sdr.get("/api/admin/analytics/sales-funnel").status_code == 403
    c = _super(db)
    assert c.get("/api/admin/analytics/sales-funnel?from=nope").status_code == 400
    assert c.get("/api/admin/analytics/sales-funnel?from=2026-03-01&to=2026-02-01").status_code == 400
    assert c.get("/api/admin/analytics/sales-funnel?from=2020-01-01&to=2026-01-01").status_code == 400
    assert c.get("/api/admin/analytics/sales-funnel?group_by=pod").status_code == 400
    assert c.get("/api/admin/analytics/sales-funnel?rep_id=missing").status_code == 404


def test_excel_export_matches_the_screen(team):
    import io
    import openpyxl
    db = team
    create_test_booking(db, create_test_lead(db, email="e@t.com", company="E", pod_id="p1"), "a", booked_at=D)
    res = _super(db).get("/api/admin/analytics/sales-funnel/export.xlsx" + Q + "&group_by=sdr")
    assert res.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(res.content))["Funnel"]
    header = [c.value for c in ws[3]]
    last = [c.value for c in ws[ws.max_row]]
    assert last[0] == "Total" and last[header.index("Meetings booked")] == 1
