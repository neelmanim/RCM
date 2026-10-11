# ── routes/export_routes.py — raw activity workbook, so the numbers on any screen can be checked row by row ─────────
# One Excel file, two sheets: every call in the range (who, which lead, owner AT THE TIME, outcome, connected /
# conversation flags that use the same SQL rules as the dashboards) and every meeting booked in the range.
# Visibility: Super Admin / Admin = everyone; Pod Admin = their own pod's people.
import io
from datetime import datetime, timedelta, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

import models
from auth import get_current_user, require_export_access
from database import get_db
from routes._admin_helpers import get_conversation_min_seconds
from routes.my_day_routes import day_window
from services import lead_lifecycle

router = APIRouter(prefix="/api/exports", tags=["Exports"])

MAX_DAYS = 92
MAX_ROWS = 50000


def _cell(v):
    """Spreadsheet-safe text: a value starting = + - @ must not run as a formula."""
    v = "" if v is None else str(v)
    return "'" + v if v[:1] in ("=", "+", "-", "@") else v


def _domain(lead):
    host = (lead.website or "").lower().replace("https://", "").replace("http://", "").replace("www.", "").split("/")[0]
    return host or ((lead.email or "").split("@")[1].lower() if "@" in (lead.email or "") else "")


def _naive(dt):
    return dt.replace(tzinfo=None) if dt is not None and dt.tzinfo else dt


@router.get("/activity.xlsx", dependencies=[Depends(require_export_access)])
def export_activity(
    date_from: str = Query(..., description="YYYY-MM-DD"),
    date_to: str = Query(..., description="YYYY-MM-DD, inclusive"),
    tz_offset_min: int = Query(0, ge=-840, le=840),
    sdr_id: Optional[str] = Query(None),
    user: dict = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    role = user.get("role")
    if role not in ("Super Admin", "Admin", "Pod Admin"):
        raise HTTPException(status_code=403, detail="Admin access required")
    start, _, _ = day_window(date_from, tz_offset_min)
    _, end, _ = day_window(date_to, tz_offset_min)
    if end <= start:
        raise HTTPException(status_code=400, detail="date_to must not be before date_from")
    if end - start > timedelta(days=MAX_DAYS):
        raise HTTPException(status_code=400, detail=f"Pick at most {MAX_DAYS} days per export")

    people = db.query(models.User)
    if role == "Pod Admin":
        if not user.get("pod_id"):
            raise HTTPException(status_code=403, detail="Your account has no pod")
        people = people.filter(models.User.pod_id == user["pod_id"])
    if sdr_id:
        people = people.filter(models.User.id == sdr_id)
    users = {u.id: u for u in people.all()}
    uids = list(users)
    min_s = get_conversation_min_seconds(db)

    rows = []   # (when, user_id, lead_id, source, outcome, duration, connected, conversation, has_notes)
    for c in (db.query(models.CallLog).filter(models.CallLog.user_id.in_(uids), models.CallLog.called_at >= start,
                                              models.CallLog.called_at < end).limit(MAX_ROWS + 1)):
        rows.append((c.called_at, c.user_id, c.lead_id, "Manual log", c.outcome, None,
                     (c.outcome or "").lower() not in models.NOT_ANSWERED_OUTCOMES_LOWER, None, bool(c.notes)))
    paired = db.query(models.CallLog.dialer_call_id).filter(models.CallLog.dialer_call_id.isnot(None))
    when = models.dialer_call_event_time()
    for d, when_v, conn, conv in (db.query(models.DialerCall, when, models.dialer_call_connected(min_s),
                                           models.dialer_call_is_conversation(min_s))
                                  .filter(models.DialerCall.user_id.in_(uids), models.DialerCall.direction == "outbound",
                                          when >= start, when < end, ~models.DialerCall.id.in_(paired))
                                  .limit(MAX_ROWS + 1)):
        rows.append((when_v, d.user_id, d.lead_id, (d.provider or "dialer").title(), d.outcome or d.status, d.duration,
                     bool(conn), bool(conv), bool(d.notes)))
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=400, detail="Too many rows — pick a shorter range or one person")
    rows.sort(key=lambda r: _naive(r[0]) or datetime.min)

    lead_ids = list({r[2] for r in rows if r[2]})
    meetings = (db.query(models.LeadMeeting).filter(models.LeadMeeting.booked_at >= start, models.LeadMeeting.booked_at < end,
                                                    models.LeadMeeting.booked_by_user_id.in_(uids)).all())
    all_ids = list(set(lead_ids) | {m.lead_id for m in meetings})
    leads = {}
    for i in range(0, len(all_ids), 500):
        for l in db.query(models.Lead).filter(models.Lead.id.in_(all_ids[i:i + 500])):
            leads[l.id] = l
    live = [l for l in leads.values() if not l.deleted_at and not getattr(l, "is_test", False)]
    states = lead_lifecycle.states_for(db, live)
    tags = {}
    for i in range(0, len(all_ids), 500):
        for lid, name in (db.query(models.lead_tags.c.lead_id, models.Tag.name)
                          .join(models.Tag, models.Tag.id == models.lead_tags.c.tag_id)
                          .filter(models.lead_tags.c.lead_id.in_(all_ids[i:i + 500]))):
            tags.setdefault(lid, []).append(name)
    batch = {l.upload_log_id for l in leads.values() if l.upload_log_id}
    batch_name = {b.id: (b.tag or b.filename or "") for b in db.query(models.LeadUploadLog).filter(
        models.LeadUploadLog.id.in_(list(batch)))} if batch else {}
    periods = {}   # lead_id -> [(valid_from, valid_to, user_id)] — only populated where the trigger has run (Postgres)
    for i in range(0, len(lead_ids), 500):
        for p in db.query(models.LeadOwnershipPeriod).filter(models.LeadOwnershipPeriod.lead_id.in_(lead_ids[i:i + 500])):
            periods.setdefault(p.lead_id, []).append((_naive(p.valid_from), _naive(p.valid_to), p.user_id))
    owner_names = dict(db.query(models.User.id, models.User.name).filter(models.User.id.in_(
        {pp[2] for v in periods.values() for pp in v}))) if periods else {}

    def owner_at(lead_id, at):
        at = _naive(at)
        for vf, vt, uid in periods.get(lead_id, []):
            if vf and at and vf <= at and (vt is None or at < vt):
                return owner_names.get(uid, "")
        return ""

    def campaign(l):
        return ", ".join(sorted(tags.get(l.id, []))) or batch_name.get(l.upload_log_id, "")

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calls"
    ws.append(["Time (UTC)", "SDR", "Lead", "Company", "Account domain", "Campaign (tags, else upload batch)",
               "Owner at the time", "Lead state now", "Status now", "Source", "Outcome", "Duration (s)",
               "Connected", "Conversation", "Has notes"])
    for at, uid, lid, src, outcome, dur, conn, conv, notes in rows:
        l = leads.get(lid)
        if not l or l.deleted_at or getattr(l, "is_test", False):
            continue
        ws.append([_naive(at).isoformat(sep=" ") if at else "", _cell(users[uid].name), _cell(f"{l.first_name or ''} {l.last_name or ''}".strip()),
                   _cell(l.company), _domain(l), _cell(campaign(l)), _cell(owner_at(lid, at)), states.get(lid, ""), l.status,
                   src, _cell(outcome), dur if dur is not None else "", "Yes" if conn else "No",
                   "" if conv is None else ("Yes" if conv else "No"), "Yes" if notes else "No"])
    ms = wb.create_sheet("Meetings")
    ms.append(["Booked at (UTC)", "Booked by", "Lead", "Company", "Account domain", "Campaign", "Meeting time (UTC)",
               "State", "Held at (UTC)", "No-show at (UTC)", "Source", "Lead status now"])
    for m in sorted(meetings, key=lambda x: _naive(x.booked_at) or datetime.min):
        l = leads.get(m.lead_id)
        if not l or l.deleted_at or getattr(l, "is_test", False):
            continue
        ms.append([_naive(m.booked_at).isoformat(sep=" ") if m.booked_at else "", _cell(users[m.booked_by_user_id].name),
                   _cell(f"{l.first_name or ''} {l.last_name or ''}".strip()), _cell(l.company), _domain(l), _cell(campaign(l)),
                   _naive(m.meeting_at).isoformat(sep=" ") if m.meeting_at else "", m.status,
                   _naive(m.held_at).isoformat(sep=" ") if m.held_at else "",
                   _naive(m.no_show_at).isoformat(sep=" ") if m.no_show_at else "", m.source or "", l.status])
    note = wb.create_sheet("Read me")
    for line in ("Calls: manual call logs + outbound dialer calls (a dialer call that was also logged by hand is counted once).",
                 "Connected / Conversation use the same rules as the dashboards (conversation threshold from Control Center).",
                 "Owner at the time is filled only for calls after ownership history began (see Data Quality); blank before that.",
                 "Campaign = lead tags, else the upload batch tag/filename. A real Campaign tag at upload is still to come.",
                 "Meetings sheet = bookings by the people in this export; bookings before v13.11 are reconstructed (source = backfill).",
                 "Lead state now = the computed lifecycle label (Unassigned / New to me / Inherited / In progress / …)."):
        note.append([line])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="activity_{date_from}_{date_to}.xlsx"'})
